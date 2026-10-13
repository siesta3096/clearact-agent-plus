"""Fixed-response runtime scenarios; these do not measure real-model task success."""

import asyncio
import hashlib
import json
import sys
from pathlib import Path

import pytest

from clearact.context.budget import ContextBudget
from clearact.context.builder import ContextBuilder
from clearact.domain.enums import RiskLevel, RunStatus
from clearact.domain.models import Action, ChatMessage, LLMResponse, Run, UserPolicy
from clearact.runtime.event_bus import EventBus
from clearact.runtime.executor import ToolExecutor
from clearact.runtime.policy import PolicyEngine
from clearact.runtime.risk import RiskEvaluator
from clearact.runtime.runner import AgentRunner
from clearact.runtime.stage_mapper import StageMapper
from clearact.storage.checkpoint_store import CheckpointStore
from clearact.storage.run_store import RunStore
from clearact.storage.snapshots import SnapshotStore
from clearact.tools.base import ToolContext
from clearact.tools.command import ExecuteCommandTool
from clearact.tools.filesystem import ReadFileTool, WriteFileTool
from clearact.tools.office import CreateDocumentTool, CreateSpreadsheetTool
from clearact.tools.registry import ToolRegistry
from clearact.tools.workflow import DeclareWorkflowPlanTool, DeclareWorkflowStepTool, ReviseWorkflowPlanTool


class ScriptedProvider:
    """Replay one fixed transcript and reject unexpected or unavailable tool calls."""

    def __init__(self, responses: list[LLMResponse]) -> None:
        self.responses = list(responses)
        self.inputs: list[list[ChatMessage]] = []

    async def chat(self, messages, tools) -> LLMResponse:
        self.inputs.append(list(messages))
        assert self.responses, "the runtime requested an unexpected model turn"
        response = self.responses.pop(0)
        offered = {tool.name for tool in tools}
        assert all(action.tool_name in offered for action in response.tool_calls)
        return response


class RecordingApprovalGate:
    def __init__(self, granted: bool) -> None:
        self.granted = granted
        self.requests: list[str] = []

    async def request(self, action, assessment) -> bool:
        assert assessment.hard_stop and assessment.category == "command_execute"
        self.requests.append(action.id)
        return self.granted


def _runner(root: Path, provider: ScriptedProvider, approval: RecordingApprovalGate) -> tuple[AgentRunner, RunStore]:
    registry = ToolRegistry()
    for tool in (
        DeclareWorkflowPlanTool(), DeclareWorkflowStepTool(), ReviseWorkflowPlanTool(),
        ReadFileTool(), WriteFileTool(), ExecuteCommandTool(), CreateDocumentTool(), CreateSpreadsheetTool(),
    ):
        registry.register(tool)
    data_root = root / "data"
    store = RunStore(data_root)
    bus = EventBus()
    bus.subscribe_sync(store.append_event)
    runner = AgentRunner(
        provider=provider, registry=registry, context_builder=ContextBuilder(ContextBudget(32_768, 0.7)),
        risk_evaluator=RiskEvaluator(root, {}), policy_engine=PolicyEngine(), approval_gate=approval,
        executor=ToolExecutor(registry, 10), event_bus=bus, run_store=store,
        checkpoint_store=CheckpointStore(data_root), tool_context=ToolContext(root, SnapshotStore(data_root, root)),
        stage_mapper=StageMapper({}), max_iterations=8, max_tool_calls=8,
    )
    return runner, store


def _plan(steps: list[dict]) -> LLMResponse:
    return LLMResponse(tool_calls=[Action(id="plan", tool_name="declare_workflow_plan", arguments={
        "summary": "执行固定场景并保留实际结果。", "steps": steps,
    })])


def _step(plan_item: dict) -> Action:
    return Action(id=f"step_{plan_item['id']}", tool_name="declare_workflow_step", arguments={
        "plan_item_id": plan_item["id"], "title": plan_item["title"],
        "summary": plan_item["summary"], "kind": plan_item["kind"],
    })


def _run(goal: str) -> Run:
    return Run(goal=goal, policy=UserPolicy(autonomy_threshold=RiskLevel.RED),
               messages=[ChatMessage(role="user", content=goal)])


def _checkpoints(root: Path, run: Run) -> list[dict]:
    path = root / "data" / "checkpoints" / f"{run.id}.jsonl"
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()] if path.exists() else []


def test_scenario_local_text_to_word_and_formula_workbook_preserves_artifacts_and_audit(workspace):
    from docx import Document
    from openpyxl import load_workbook

    source = "恢复文件: 8\n断连续接: 9\n"
    (workspace / "samples.txt").write_text(source, encoding="utf-8")
    phases = [
        {"id": "source", "title": "读取样本", "summary": "读取本地固定样本。", "kind": "research"},
        {"id": "artifacts", "title": "生成交付物", "summary": "生成报告和统计工作簿。", "kind": "files"},
    ]
    read = Action(id="read_samples", tool_name="read_file", arguments={"path": "samples.txt"})
    word = Action(id="create_report", tool_name="create_document", arguments={
        "path": "reports/summary.docx", "title": "恢复场景统计", "sections": [{
            "heading": "样本结果", "paragraphs": ["固定样本包含恢复文件 8 次、断连续接 9 次。"],
            "table": {"headers": ["场景", "次数"], "rows": [["恢复文件", "8"], ["断连续接", "9"]]},
        }],
    })
    excel = Action(id="create_counts", tool_name="create_spreadsheet", arguments={
        "path": "reports/counts.xlsx", "sheets": [{"name": "Counts", "headers": ["场景", "次数"],
            "rows": [["恢复文件", 8], ["断连续接", 9], ["合计", "=SUM(B2:B3)"]],
            "column_formats": ["text", "integer"]}],
    })
    provider = ScriptedProvider([
        _plan(phases), LLMResponse(tool_calls=[_step(phases[0]), read]),
        LLMResponse(tool_calls=[_step(phases[1]), word, excel]),
        LLMResponse(content="已生成 Word 和 Excel；工作簿公式将在 Excel 打开时重新计算。"),
    ])
    approval = RecordingApprovalGate(False)
    runner, store = _runner(workspace, provider, approval)
    run = _run("只根据本地 samples.txt 生成 reports/summary.docx 和 reports/counts.xlsx，保留汇总公式。")
    asyncio.run(runner.run(run))

    document = Document(workspace / "reports" / "summary.docx")
    assert document.paragraphs[0].text == "恢复场景统计"
    assert "恢复文件 8 次、断连续接 9 次" in document.paragraphs[2].text
    assert document.tables[0].cell(2, 1).text == "9"
    workbook = load_workbook(workspace / "reports" / "counts.xlsx", data_only=False)
    assert workbook["Counts"]["B2"].value == 8
    assert workbook["Counts"]["B3"].value == 9
    assert workbook["Counts"]["B4"].value == "=SUM(B2:B3)"
    assert workbook["Counts"]["B4"].data_type == "f"
    workbook.close()

    saved = store.load_run(run.id)
    results = {message.tool_call_id: message for message in saved.messages if message.role == "tool"}
    assert saved.status is RunStatus.COMPLETED
    assert results[read.id].content == source
    assert all(results[action.id].metadata["status"] == "succeeded" for action in (read, word, excel))
    assert results[excel.id].metadata["formula_count"] == 1
    assert results[excel.id].metadata["formulas_calculated"] is False
    for action in (word, excel):
        metadata = results[action.id].metadata
        assert metadata["snapshot_id"]
        assert metadata["content_sha256"] == hashlib.sha256(Path(metadata["path"]).read_bytes()).hexdigest()
    assert [checkpoint["action_id"] for checkpoint in _checkpoints(workspace, run)] == [read.id, word.id, excel.id]
    assert [event.action_id for event in store.load_events(run.id) if event.type == "action.completed"] == [
        read.id, word.id, excel.id,
    ]
    assert all(step.status == "completed" for step in saved.workflow_steps)
    assert any(message.tool_call_id == read.id and message.content == source for message in provider.inputs[2])
    assert provider.responses == [] and approval.requests == []


def test_scenario_declined_command_never_creates_marker_and_records_denial(workspace):
    phase = {"id": "command", "title": "执行验证命令", "summary": "确认后运行固定验证命令。", "kind": "external"}
    command = Action(id="declined_command", tool_name="execute_command", arguments={"command": [
        sys.executable, "-c", "from pathlib import Path; Path('marker.txt').write_text('executed', encoding='utf-8')",
    ]})
    provider = ScriptedProvider([
        _plan([phase]), LLMResponse(tool_calls=[_step(phase), command]),
        LLMResponse(content="命令未获批准，未执行。"),
    ])
    approval = RecordingApprovalGate(False)
    runner, store = _runner(workspace, provider, approval)
    run = _run("请求确认后执行 Python 验证命令。")
    asyncio.run(runner.run(run))

    saved = store.load_run(run.id)
    result = next(message for message in saved.messages if message.tool_call_id == command.id)
    assert not (workspace / "marker.txt").exists()
    assert approval.requests == [command.id]
    assert result.metadata == {"status": "denied", "reason": "user_declined"}
    assert [event.type for event in store.load_events(run.id) if event.action_id == command.id] == [
        "approval.required", "action.denied",
    ]
    assert _checkpoints(workspace, run) == []
    assert saved.workflow_steps[-1].status == "needs_review"
    assert provider.responses == []


@pytest.mark.parametrize("exit_code", [0, 7], ids=["successful-exit", "failed-exit"])
def test_scenario_approved_command_records_output_and_only_checkpoints_success(workspace, exit_code):
    phase = {"id": "command", "title": "执行验证命令", "summary": "确认后捕获固定输出和退出码。", "kind": "external"}
    command = Action(id="approved_command", tool_name="execute_command", arguments={"command": [
        sys.executable, "-c",
        f"import sys; print('scenario-output'); print('scenario-stderr', file=sys.stderr); sys.exit({exit_code})",
    ], "timeout_seconds": 5})
    provider = ScriptedProvider([
        _plan([phase]), LLMResponse(tool_calls=[_step(phase), command]),
        LLMResponse(content="验证命令已完成。" if exit_code == 0 else "验证命令失败，退出码为 7。"),
    ])
    approval = RecordingApprovalGate(True)
    runner, store = _runner(workspace, provider, approval)
    run = _run("请求确认后执行 Python 验证命令并报告 stdout、stderr 和退出码。")
    asyncio.run(runner.run(run))

    saved = store.load_run(run.id)
    result = next(message for message in saved.messages if message.tool_call_id == command.id)
    assert approval.requests == [command.id]
    assert result.metadata["stdout"].strip() == "scenario-output"
    assert result.metadata["stderr"].strip() == "scenario-stderr"
    assert result.metadata["exit_code"] == exit_code
    assert result.metadata["timed_out"] is False
    assert result.metadata["status"] == ("succeeded" if exit_code == 0 else "failed")
    assert [checkpoint["action_id"] for checkpoint in _checkpoints(workspace, run)] == (
        [command.id] if exit_code == 0 else []
    )
    assert [event.type for event in store.load_events(run.id) if event.action_id == command.id] == [
        "approval.required", "action.started", "action.completed" if exit_code == 0 else "action.failed",
    ]
    assert saved.workflow_steps[-1].status == ("completed" if exit_code == 0 else "needs_review")
    assert provider.responses == []

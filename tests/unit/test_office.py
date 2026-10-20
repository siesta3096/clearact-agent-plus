import asyncio
import builtins
import os
from datetime import datetime
from io import BytesIO
from threading import Event

import pytest

import clearact.tools.artifacts as artifacts
import clearact.tools.office as office
from clearact.domain.errors import ScopeViolationError, ToolValidationError
from clearact.storage.snapshots import SnapshotStore
from clearact.tools.base import ToolContext
from clearact.tools.office import CreateDocumentTool, CreateSpreadsheetTool


def _document(**overrides):
    return {
        "path": "reports/报告.docx",
        "title": "课程项目报告",
        "subtitle": "2026 年秋季",
        "sections": [
            {
                "heading": "执行与恢复",
                "paragraphs": ["程序保留用户后来修改的文件。"],
                "bullets": ["建立快照", "检查文件内容"],
                "table": {"headers": ["场景", "结果"], "rows": [["用户修改", "保留"]]},
            }
        ],
        **overrides,
    }


def _spreadsheet(**overrides):
    return {
        "path": "reports/统计.xlsx",
        "sheets": [
            {
                "name": "运行统计",
                "headers": ["日期", "任务", "成功数", "比率", "汇总", "原始文本"],
                "rows": [
                    ["2026-10-13", "文件恢复", 8, 0.8, "=SUM(C2:C3)", "=literal"],
                    ["2026-10-14", "断连续接", 9, 0.9, None, "#N/A"],
                ],
                "column_formats": ["date", "general", "integer", "percent", "integer", "text"],
                "column_widths": [16, None, 12, 12, 14, 16],
                "charts": [
                    {"type": "bar", "title": "成功数", "category_column": 2, "value_columns": [3], "anchor": "H2"},
                    {"type": "line", "category_column": 2, "value_columns": [4], "anchor": "H20"},
                ],
            }
        ],
        **overrides,
    }


def test_document_contains_native_content_layout_and_page_numbers(workspace):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm

    result = asyncio.run(CreateDocumentTool().execute(_document(), ToolContext(workspace), "word"))
    target = workspace / "reports" / "报告.docx"
    document = Document(BytesIO(target.read_bytes()))

    assert result.ok
    assert result.metadata["path"] == str(target)
    assert result.metadata["size_bytes"] == target.stat().st_size
    assert result.metadata["file_kind"] == "docx"
    assert result.metadata["section_count"] == result.metadata["table_count"] == 1
    assert [paragraph.text for paragraph in document.paragraphs][:4] == [
        "课程项目报告",
        "2026 年秋季",
        "执行与恢复",
        "程序保留用户后来修改的文件。",
    ]
    assert document.paragraphs[4].style.name == "List Bullet"
    assert document.tables[0].cell(1, 1).text == "保留"
    assert document.tables[0].rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
    assert document.sections[0].page_width == pytest.approx(Mm(210), abs=1000)
    assert document.sections[0].page_height == pytest.approx(Mm(297), abs=1000)
    assert document.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")) == "微软雅黑"
    assert document.sections[0].header.paragraphs[0].text == "课程项目报告"
    assert 'w:instr="PAGE"' in document.sections[0].footer._element.xml


@pytest.mark.parametrize("number_format", ["number", "Number", "NUMBER"])
def test_number_format_alias_roundtrips_numbers_instead_of_dates(workspace, number_format):
    from openpyxl import load_workbook

    result = asyncio.run(CreateSpreadsheetTool().execute({
        "path": "sales.xlsx", "sheets": [{"name": "Sales", "headers": ["category", "amount"],
            "rows": [["A", 12.5], ["B", 7], ["A", 2.5], ["TOTAL", "=SUM(B2:B4)"]],
            "column_formats": ["text", number_format]}],
    }, ToolContext(workspace), "excel"))
    assert result.ok
    workbook = load_workbook(workspace / "sales.xlsx", data_only=False)
    try:
        sheet = workbook["Sales"]
        assert [sheet[f"B{row}"].value for row in range(2, 5)] == [12.5, 7, 2.5]
        assert sheet["B2"].number_format == "General"
        assert sheet["B5"].value == "=SUM(B2:B4)"
    finally:
        workbook.close()


def test_spreadsheet_preserves_types_formats_formulas_and_native_charts(workspace):
    from openpyxl import load_workbook

    result = asyncio.run(CreateSpreadsheetTool().execute(_spreadsheet(), ToolContext(workspace), "excel"))
    target = workspace / "reports" / "统计.xlsx"
    workbook = load_workbook(target, data_only=False)
    sheet = workbook["运行统计"]

    assert result.ok
    assert result.metadata["sheets"] == [{"name": "运行统计", "rows": 2, "columns": 6}]
    assert result.metadata["chart_count"] == 2
    assert result.metadata["formula_count"] == 1
    assert result.metadata["formulas_calculated"] is False
    assert result.metadata["recalculate_on_open"] is True
    assert sheet["A2"].value == datetime(2026, 10, 13)
    assert sheet["A2"].number_format == "yyyy-mm-dd"
    assert sheet["C2"].value == 8
    assert sheet["D2"].number_format == "0.00%"
    assert sheet["E2"].value == "=SUM(C2:C3)" and sheet["E2"].data_type == "f"
    assert sheet["F2"].value == "=literal" and sheet["F2"].data_type == "s"
    assert sheet["F3"].value == "#N/A" and sheet["F3"].data_type == "s"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:F3"
    assert sheet.tables["ClearActTable1"].tableStyleInfo.showRowStripes is True
    assert sheet.column_dimensions["A"].width == 16
    assert len(sheet._charts) == 2
    assert workbook.calculation.calcMode == "auto"
    assert workbook.calculation.fullCalcOnLoad and workbook.calculation.forceFullCalc
    workbook.close()
    cached = load_workbook(target, data_only=True)
    assert cached["运行统计"]["E2"].value is None
    cached.close()


@pytest.mark.parametrize(
    "tool, arguments",
    [
        (CreateDocumentTool(), _document(path="wrong.pdf")),
        (CreateDocumentTool(), _document(sections=[{"paragraphs": ["valid"]}, {"heading": []}])),
        (CreateDocumentTool(), _document(sections=[{"table": {"headers": ["A"], "rows": [["B", "C"]]}}])),
        (CreateDocumentTool(), _document(title="invalid\x00text")),
        (CreateSpreadsheetTool(), _spreadsheet(path="wrong.xls")),
        (CreateSpreadsheetTool(), _spreadsheet(sheets=[{"name": "Bad/Name", "headers": ["A"], "rows": []}])),
        (CreateSpreadsheetTool(), _spreadsheet(sheets=[{"name": "Data", "headers": ["A", "a"], "rows": []}])),
        (CreateSpreadsheetTool(), _spreadsheet(sheets=[{"name": "Data", "headers": ["A"], "rows": [[float("nan")]]}])),
        (CreateSpreadsheetTool(), _spreadsheet(sheets=[{"name": "Data", "headers": ["A"], "rows": [["x\x01"]]}])),
        (
            CreateSpreadsheetTool(),
            _spreadsheet(
                sheets=[
                    {
                        "name": "Data",
                        "headers": ["A"],
                        "rows": [["2026-10-13"]],
                        "column_formats": ["datetime"],
                        "column_widths": [float("inf")],
                    }
                ]
            ),
        ),
        (
            CreateSpreadsheetTool(),
            _spreadsheet(
                sheets=[{"name": "Data", "headers": ["A"], "rows": [["yesterday"]], "column_formats": ["date"]}]
            ),
        ),
        (
            CreateSpreadsheetTool(),
            _spreadsheet(
                sheets=[
                    {"name": "Data", "headers": ["A"], "rows": [[1]], "charts": [{"type": [], "value_columns": [1]}]}
                ]
            ),
        ),
    ],
)
def test_invalid_input_does_not_create_files_or_directories(workspace, tool, arguments):
    with pytest.raises(ToolValidationError):
        asyncio.run(tool.execute(arguments, ToolContext(workspace), "invalid"))
    assert list(workspace.iterdir()) == []


def test_invalid_later_sheet_does_not_replace_existing_workbook(workspace):
    target = workspace / "existing.xlsx"
    target.write_bytes(b"original file")
    arguments = _spreadsheet(path="existing.xlsx")
    arguments["sheets"].append({"name": "运行统计", "headers": ["A"], "rows": [[1]]})

    with pytest.raises(ToolValidationError, match="unique"):
        asyncio.run(CreateSpreadsheetTool().execute(arguments, ToolContext(workspace), "invalid"))
    assert target.read_bytes() == b"original file"


@pytest.mark.parametrize(
    "tool, arguments",
    [
        (CreateDocumentTool(), _document(path="../outside.docx")),
        (CreateSpreadsheetTool(), _spreadsheet(path="../outside.xlsx")),
    ],
)
def test_office_files_respect_workspace_scope(workspace, tool, arguments):
    with pytest.raises(ScopeViolationError):
        asyncio.run(tool.execute(arguments, ToolContext(workspace), "escape"))


@pytest.mark.parametrize(
    "tool, arguments",
    [
        (CreateDocumentTool(), _document(path="report.docx")),
        (CreateSpreadsheetTool(), _spreadsheet(path="report.xlsx")),
    ],
)
def test_overwritten_office_files_can_restore_original_bytes(workspace, tmp_path, tool, arguments):
    context = ToolContext(workspace, SnapshotStore(tmp_path / "data", workspace))
    first = asyncio.run(tool.execute(arguments, context, "first"))
    target = workspace / arguments["path"]
    original = target.read_bytes()
    if tool.name == "create_document":
        arguments["title"] = "更新报告"
    else:
        arguments["sheets"][0]["rows"][0][2] = 42
    second = asyncio.run(tool.execute(arguments, context, "second"))
    assert second.metadata["snapshot_id"]
    assert target.read_bytes() != original
    context.snapshot_store.restore(second.metadata["snapshot_id"])
    assert target.read_bytes() == original
    context.snapshot_store.restore(first.metadata["snapshot_id"])
    assert not target.exists()


@pytest.mark.parametrize(
    "tool, arguments, package",
    [
        (CreateDocumentTool(), _document(), "docx"),
        (CreateSpreadsheetTool(), _spreadsheet(), "openpyxl"),
    ],
)
def test_missing_office_dependency_reports_clear_error(workspace, monkeypatch, tool, arguments, package):
    original_import = builtins.__import__

    def guarded_import(name, *args, **kwargs):
        if name.split(".")[0] == package:
            raise ImportError("dependency unavailable")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", guarded_import)
    with pytest.raises(ToolValidationError, match="generation requires"):
        asyncio.run(tool.execute(arguments, ToolContext(workspace), "missing"))
    assert list(workspace.iterdir()) == []


def test_spreadsheet_bound_limits_are_checked_before_writing(workspace):
    arguments = _spreadsheet(
        sheets=[{"name": "TooBig", "headers": [str(i) for i in range(50)], "rows": [[0] * 50 for _ in range(2000)]}]
    )
    with pytest.raises(ToolValidationError, match="cell limit"):
        asyncio.run(CreateSpreadsheetTool().execute(arguments, ToolContext(workspace), "too-big"))
    assert list(workspace.iterdir()) == []


def test_empty_table_multiple_sheets_and_literal_headers(workspace):
    from openpyxl import load_workbook

    arguments = _spreadsheet(
        sheets=[
            {"name": "Empty", "headers": ["=literal header"], "rows": []},
            {"name": "Other", "headers": ["Value"], "rows": [[True], [None]]},
        ]
    )
    result = asyncio.run(CreateSpreadsheetTool().execute(arguments, ToolContext(workspace), "multiple"))
    workbook = load_workbook(result.metadata["path"])
    assert workbook.sheetnames == ["Empty", "Other"]
    assert workbook["Empty"]["A1"].data_type == "s"
    assert workbook["Empty"].tables["ClearActTable1"].ref == "A1:A1"
    assert workbook["Other"].tables["ClearActTable2"].ref == "A1:A3"
    assert workbook["Other"]["A2"].value is True
    workbook.close()


def test_default_chart_anchors_do_not_overlap(workspace):
    from openpyxl import load_workbook

    arguments = _spreadsheet()
    for chart in arguments["sheets"][0]["charts"]:
        chart.pop("anchor")
    result = asyncio.run(CreateSpreadsheetTool().execute(arguments, ToolContext(workspace), "charts"))
    workbook = load_workbook(result.metadata["path"])
    charts = workbook["运行统计"]._charts
    assert charts[0].anchor._from.row == 5
    assert charts[1].anchor._from.row == 23
    workbook.close()


@pytest.mark.parametrize(
    "tool, arguments, renderer",
    [
        (CreateDocumentTool(), _document(), "_render_document"),
        (CreateSpreadsheetTool(), _spreadsheet(), "_render_spreadsheet"),
    ],
)
def test_cancelled_generation_never_commits_after_renderer_finishes(workspace, monkeypatch, tool, arguments, renderer):
    started, release, finished = Event(), Event(), Event()
    writes = []
    original_save = artifacts.save_artifact

    def paused_renderer(_content):
        started.set()
        try:
            assert release.wait(5), "test must release the renderer thread"
            return b"rendered artifact bytes"
        finally:
            finished.set()

    def tracked_save(*args, **kwargs):
        writes.append(args[1])
        return original_save(*args, **kwargs)

    monkeypatch.setattr(office, renderer, paused_renderer)
    monkeypatch.setattr(artifacts, "save_artifact", tracked_save)

    async def cancel_during_generation():
        task = asyncio.create_task(tool.execute(arguments, ToolContext(workspace), "cancel"))
        try:
            assert await asyncio.to_thread(started.wait, 5)
            task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await task
        finally:
            release.set()
        assert await asyncio.to_thread(finished.wait, 5)

    asyncio.run(cancel_during_generation())
    assert writes == []
    assert list(workspace.iterdir()) == []


@pytest.mark.parametrize("existed", [False, True])
def test_spreadsheet_preserves_user_changes_during_rendering_without_snapshot(
    workspace, tmp_path, monkeypatch, existed,
):
    target = workspace / "report.xlsx"
    if existed:
        target.write_bytes(b"Previously approved workbook")
    user_edit = b"User's updated workbook"
    data_root = tmp_path / "data"
    context = ToolContext(workspace, SnapshotStore(data_root, workspace))

    def render_with_user_edit(_content):
        target.write_bytes(user_edit)
        return b"Generated workbook"

    monkeypatch.setattr(office, "_render_spreadsheet", render_with_user_edit)
    with pytest.raises(ToolValidationError, match="changed"):
        asyncio.run(CreateSpreadsheetTool().execute(_spreadsheet(path="report.xlsx"), context, "concurrent"))
    assert target.read_bytes() == user_edit
    assert list((data_root / "snapshots").iterdir()) == []


@pytest.mark.parametrize(
    "tool, arguments, renderer",
    [
        (CreateDocumentTool(), _document(path="report.docx"), "_render_document"),
        (CreateSpreadsheetTool(), _spreadsheet(path="report.xlsx"), "_render_spreadsheet"),
    ],
)
def test_office_rechecks_changed_path_resolution_before_commit(
    workspace, tmp_path, monkeypatch, tool, arguments, renderer,
):
    original_target = workspace / arguments["path"]
    original_target.write_bytes(b"Original approved file")
    redirected_target = workspace / "redirected" / arguments["path"]
    redirected_target.parent.mkdir()
    redirected_target.write_bytes(b"User's file in redirected directory")
    data_root = tmp_path / "data"
    context = ToolContext(workspace, SnapshotStore(data_root, workspace))
    original_resolve = tool._resolve
    redirected = False

    def changed_resolution(context, raw_path, *, writing=False):
        # Simulate a directory symlink being redirected without requiring Windows symlink privileges.
        return redirected_target if redirected else original_resolve(context, raw_path, writing=writing)

    def redirect_during_render(_content):
        nonlocal redirected
        redirected = True
        return b"Generated artifact"

    monkeypatch.setattr(tool, "_resolve", changed_resolution)
    monkeypatch.setattr(office, renderer, redirect_during_render)
    with pytest.raises(ToolValidationError, match="path changed"):
        asyncio.run(tool.execute(arguments, context, "redirected"))
    assert original_target.read_bytes() == b"Original approved file"
    assert redirected_target.read_bytes() == b"User's file in redirected directory"
    assert list((data_root / "snapshots").iterdir()) == []


def test_artifact_version_preserves_same_length_edit_with_restored_timestamp(workspace, tmp_path):
    target = workspace / "report.docx"
    target.write_bytes(b"Original bytes")
    initial_stat = target.stat()
    expected = artifacts.capture_artifact_version(target)
    target.write_bytes(b"Modified bytes")
    os.utime(target, ns=(initial_stat.st_atime_ns, initial_stat.st_mtime_ns))
    data_root = tmp_path / "data"
    context = ToolContext(workspace, SnapshotStore(data_root, workspace))

    with pytest.raises(ToolValidationError, match="changed"):
        artifacts.save_artifact(context, target, b"Generated bytes", expected_version=expected)
    assert target.read_bytes() == b"Modified bytes"
    assert list((data_root / "snapshots").iterdir()) == []


def test_artifact_save_keeps_optional_version_check_compatible(workspace):
    target = workspace / "report.xlsx"
    target.write_bytes(b"Original bytes")
    result = artifacts.save_artifact(ToolContext(workspace), target, b"Replacement bytes")
    assert target.read_bytes() == b"Replacement bytes"
    assert result["existed"] is True

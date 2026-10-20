from __future__ import annotations

import asyncio
import math
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from clearact.domain.errors import ToolValidationError
from clearact.domain.models import ToolDefinition, ToolResult
from clearact.tools.artifacts import ArtifactVersion, capture_artifact_version
from clearact.tools.base import ToolContext
from clearact.tools.filesystem import _FilesystemTool

MAX_TOTAL_TEXT = 500_000
MAX_CELLS = 100_000
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_NUMBER_FORMATS = {
    "general": "General",
    "number": "General",
    "text": "@",
    "integer": "#,##0",
    "decimal": "#,##0.00",
    "percent": "0.00%",
    "currency": '"¥"#,##0.00',
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm:ss",
}


def _object(value: Any, label: str, allowed: set[str]) -> dict:
    if not isinstance(value, dict) or set(value) - allowed:
        raise ToolValidationError(f"{label} must be an object with supported fields: {', '.join(sorted(allowed))}")
    return value


def _text(value: Any, label: str, *, limit: int = 20_000, nonempty: bool = False) -> str:
    if not isinstance(value, str) or len(value) > limit or _INVALID_XML.search(value):
        raise ToolValidationError(f"{label} must be valid text of at most {limit} characters")
    if nonempty and not value.strip():
        raise ToolValidationError(f"{label} must not be empty")
    return value


def _list(value: Any, label: str, *, maximum: int, minimum: int = 0) -> list:
    if not isinstance(value, list) or not minimum <= len(value) <= maximum:
        raise ToolValidationError(f"{label} must be an array with {minimum} to {maximum} items")
    return value


def _document_content(arguments: dict) -> dict:
    _object(arguments, "arguments", {"path", "title", "subtitle", "sections"})
    title = _text(arguments.get("title"), "title", limit=300, nonempty=True)
    subtitle = _text(arguments.get("subtitle", ""), "subtitle", limit=1000)
    sections = _list(arguments.get("sections"), "sections", minimum=1, maximum=80)
    total_text, total_cells = len(title) + len(subtitle), 0
    normalized = []
    for index, raw_section in enumerate(sections, 1):
        label = f"sections[{index}]"
        section = _object(raw_section, label, {"heading", "paragraphs", "bullets", "table"})
        heading = _text(section.get("heading", ""), f"{label}.heading", limit=300)
        item = {"heading": heading}
        total_text += len(heading)
        for field in ("paragraphs", "bullets"):
            texts = _list(section.get(field, []), f"{label}.{field}", maximum=200)
            item[field] = [_text(text, f"{label}.{field}") for text in texts]
            total_text += sum(map(len, item[field]))
        if "table" in section:
            table = _object(section["table"], f"{label}.table", {"headers", "rows"})
            headers = _list(table.get("headers"), f"{label}.table.headers", minimum=1, maximum=12)
            headers = [_text(header, "table header", limit=300, nonempty=True) for header in headers]
            rows = _list(table.get("rows"), f"{label}.table.rows", maximum=500)
            normalized_rows = []
            for row in rows:
                _list(row, "table row", minimum=len(headers), maximum=len(headers))
                normalized_rows.append([_text(cell, "table cell", limit=2000) for cell in row])
            total_cells += len(headers) * (len(rows) + 1)
            total_text += sum(map(len, headers)) + sum(len(cell) for row in normalized_rows for cell in row)
            item["table"] = {"headers": headers, "rows": normalized_rows}
        if not (heading or item["paragraphs"] or item["bullets"] or "table" in item):
            raise ToolValidationError(f"{label} must contain some content")
        if total_text > MAX_TOTAL_TEXT or total_cells > MAX_CELLS:
            raise ToolValidationError("document exceeds the total text or table cell limit")
        normalized.append(item)
    return {"title": title, "subtitle": subtitle, "sections": normalized}


def _cell(value: Any, label: str, number_format: str) -> Any:
    if value is None or type(value) is bool:
        return value
    if type(value) in {int, float}:
        if abs(value) > 1e308 or (type(value) is float and not math.isfinite(value)):
            raise ToolValidationError(f"{label} must be a finite Excel number")
        return value
    text = _text(value, label, limit=32_767)
    if not text or text.startswith("=") or number_format not in {"date", "datetime"}:
        return text
    try:
        parsed = date.fromisoformat(text) if number_format == "date" else datetime.fromisoformat(text)
    except ValueError as exc:
        raise ToolValidationError(f"{label} must be an ISO {number_format} for the selected format") from exc
    if isinstance(parsed, datetime) and parsed.tzinfo is not None:
        raise ToolValidationError(f"{label} must be a timezone-free datetime for Excel")
    if parsed.year < 1900:
        raise ToolValidationError(f"{label} must be a date from 1900 onward")
    return parsed


def _spreadsheet_content(arguments: dict) -> list[dict]:
    _object(arguments, "arguments", {"path", "sheets"})
    sheets = _list(arguments.get("sheets"), "sheets", minimum=1, maximum=12)
    names: set[str] = set()
    total_text, total_cells = 0, 0
    normalized = []
    for raw_sheet in sheets:
        sheet = _object(raw_sheet, "sheet", {"name", "headers", "rows", "column_formats", "column_widths", "charts"})
        name = _text(sheet.get("name"), "sheet.name", limit=31, nonempty=True)
        if re.search(r"[\\/*?:\[\]]", name) or name.startswith("'") or name.endswith("'"):
            raise ToolValidationError("sheet.name contains characters forbidden by Excel")
        if name.casefold() in names:
            raise ToolValidationError("sheet names must be unique, ignoring case")
        names.add(name.casefold())
        headers = _list(sheet.get("headers"), "sheet.headers", minimum=1, maximum=50)
        headers = [_text(header, "sheet header", limit=255, nonempty=True) for header in headers]
        if len({header.casefold() for header in headers}) != len(headers):
            raise ToolValidationError("sheet headers must be unique, ignoring case")
        count = len(headers)
        formats = _list(
            sheet.get("column_formats", ["general"] * count), "column_formats", minimum=count, maximum=count
        )
        formats = [_text(fmt, "column format", limit=128, nonempty=True) for fmt in formats]
        formats = [fmt.casefold() if fmt.casefold() in _NUMBER_FORMATS else fmt for fmt in formats]
        widths = _list(sheet.get("column_widths", [None] * count), "column_widths", minimum=count, maximum=count)
        if any(width is not None and (type(width) not in {int, float} or not 5 <= width <= 80) for width in widths):
            raise ToolValidationError("column widths must be null (automatic) or numbers from 5 to 80")
        rows = _list(sheet.get("rows"), "sheet.rows", maximum=10_000)
        total_cells += count * (len(rows) + 1)
        if total_cells > MAX_CELLS:
            raise ToolValidationError("spreadsheet exceeds the total text or cell limit")
        normalized_rows = []
        for row_index, row in enumerate(rows, 2):
            _list(row, f"row {row_index}", minimum=count, maximum=count)
            normalized_row = [
                _cell(value, f"row {row_index} column {column + 1}", formats[column])
                for column, value in enumerate(row)
            ]
            total_text += sum(len(value) for value in normalized_row if isinstance(value, str))
            if total_text > MAX_TOTAL_TEXT:
                raise ToolValidationError("spreadsheet exceeds the total text or cell limit")
            normalized_rows.append(normalized_row)
        total_text += len(name) + sum(map(len, headers)) + sum(map(len, formats))
        charts = _list(sheet.get("charts", []), "charts", maximum=4)
        normalized_charts = []
        for chart_index, chart in enumerate(charts):
            _object(chart, "chart", {"type", "title", "category_column", "value_columns", "anchor"})
            chart_type = chart.get("type")
            if not isinstance(chart_type, str) or chart_type not in {"bar", "line"}:
                raise ToolValidationError("chart.type must be bar or line")
            category = chart.get("category_column", 1)
            values = _list(chart.get("value_columns"), "chart.value_columns", minimum=1, maximum=count)
            if (
                type(category) is not int
                or not 1 <= category <= count
                or any(type(value) is not int or not 1 <= value <= count for value in values)
                or len(set(values)) != len(values)
                or category in values
            ):
                raise ToolValidationError("chart columns must be distinct valid 1-based column numbers")
            if not rows:
                raise ToolValidationError("charts require at least one data row")
            title = _text(chart.get("title", ""), "chart.title", limit=300)
            anchor = _text(chart.get("anchor", f"A{len(rows) + 4 + chart_index * 18}"), "chart.anchor", limit=12)
            match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", anchor)
            if not match or int(match[2]) > 1_048_576 or _column_index(match[1]) > 16_384:
                raise ToolValidationError("chart.anchor must be a valid Excel cell address, such as A15")
            total_text += len(title)
            normalized_charts.append(
                {
                    "type": chart_type,
                    "title": title,
                    "category_column": category,
                    "value_columns": values,
                    "anchor": anchor,
                }
            )
        if total_text > MAX_TOTAL_TEXT or total_cells > MAX_CELLS:
            raise ToolValidationError("spreadsheet exceeds the total text or cell limit")
        normalized.append(
            {
                "name": name,
                "headers": headers,
                "rows": normalized_rows,
                "column_formats": formats,
                "column_widths": widths,
                "charts": normalized_charts,
            }
        )
    return normalized


def _column_index(letters: str) -> int:
    result = 0
    for letter in letters:
        result = result * 26 + ord(letter) - ord("A") + 1
    return result


def _render_document(content: dict) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
    except ImportError as exc:
        raise ToolValidationError("Word generation requires python-docx; reinstall ClearAct's dependencies") from exc

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(22)
    section.left_margin = section.right_margin = Mm(22)
    section.header_distance = section.footer_distance = Mm(10)
    for name, size in (
        ("Normal", 11),
        ("Title", 24),
        ("Subtitle", 12),
        ("Heading 1", 15),
        ("List Bullet", 11),
        ("Header", 9),
        ("Footer", 9),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
        style.paragraph_format.widow_control = True
        style.paragraph_format.space_after = Pt(6)
    document.styles["Normal"].paragraph_format.line_spacing = 1.15
    document.styles["Heading 1"].font.color.rgb = RGBColor.from_string("17365D")
    document.styles["Heading 1"].paragraph_format.keep_with_next = True
    document.add_paragraph(content["title"], "Title")
    if content["subtitle"]:
        document.add_paragraph(content["subtitle"], "Subtitle")
    section.header.paragraphs[0].text = content["title"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    for item in content["sections"]:
        if item["heading"]:
            document.add_heading(item["heading"], level=1)
        for paragraph in item["paragraphs"]:
            document.add_paragraph(paragraph)
        for bullet in item["bullets"]:
            document.add_paragraph(bullet, "List Bullet")
        if "table" not in item:
            continue
        table_content = item["table"]
        table = document.add_table(rows=1, cols=len(table_content["headers"]))
        table.style = "Table Grid"
        table.autofit = False
        for column in table.columns:
            column.width = Mm(166 / len(table_content["headers"]))
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, header in zip(table.rows[0].cells, table_content["headers"], strict=True):
            cell.text = header
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E8EFF7")
            cell._tc.get_or_add_tcPr().append(shading)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in table_content["rows"]:
            cells = table.add_row().cells
            for cell, text in zip(cells, row, strict=True):
                cell.text = text
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)
        document.add_paragraph()
    document.core_properties.title = content["title"]
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _display_width(value: Any) -> int:
    return sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in str(value or "")[:100])


def _render_spreadsheet(sheets: list[dict]) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.workbook.properties import CalcProperties
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise ToolValidationError("Excel generation requires openpyxl; reinstall ClearAct's dependencies") from exc

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    for index, sheet in enumerate(sheets, 1):
        worksheet = workbook.create_sheet(sheet["name"])
        worksheet.append(sheet["headers"])
        for row in sheet["rows"]:
            worksheet.append(row)
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        worksheet.print_title_rows = "1:1"
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.oddFooter.center.text = "Page &P of &N"
        for cell in worksheet[1]:
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.data_type = "s"
        worksheet.row_dimensions[1].height = 28
        for column_index, number_format in enumerate(sheet["column_formats"], 1):
            for row_index in range(2, len(sheet["rows"]) + 2):
                cell = worksheet.cell(row_index, column_index)
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.number_format = _NUMBER_FORMATS.get(number_format, number_format)
                if isinstance(cell.value, str) and (number_format == "text" or not cell.value.startswith("=")):
                    cell.data_type = "s"
            width = sheet["column_widths"][column_index - 1]
            if width is None:
                sample = [sheet["headers"][column_index - 1]] + [row[column_index - 1] for row in sheet["rows"][:200]]
                width = max(10, min(45, max(map(_display_width, sample)) + 2))
            worksheet.column_dimensions[get_column_letter(column_index)].width = width
        last_column = get_column_letter(len(sheet["headers"]))
        table_range = f"A1:{last_column}{len(sheet['rows']) + 1}"
        table = Table(displayName=f"ClearActTable{index}", ref=table_range)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        worksheet.add_table(table)
        worksheet.auto_filter.ref = table_range
        for spec in sheet["charts"]:
            chart = BarChart() if spec["type"] == "bar" else LineChart()
            if spec["title"]:
                chart.title = spec["title"]
            chart.style = 10
            for column in spec["value_columns"]:
                chart.add_data(
                    Reference(worksheet, min_col=column, max_col=column, min_row=1, max_row=len(sheet["rows"]) + 1),
                    titles_from_data=True,
                )
            chart.set_categories(
                Reference(worksheet, min_col=spec["category_column"], min_row=2, max_row=len(sheet["rows"]) + 1)
            )
            worksheet.add_chart(chart, spec["anchor"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class _OfficeTool(_FilesystemTool):
    suffix: str

    def _target(self, arguments: dict, context: ToolContext) -> Path:
        raw_path = _text(arguments.get("path"), "path", limit=4096, nonempty=True)
        target = self._resolve(context, raw_path, writing=True)
        if target.suffix.lower() != self.suffix:
            raise ToolValidationError(f"{self.name} requires a {self.suffix} path")
        if target.exists() and not target.is_file():
            raise ToolValidationError("path must name a file, not a directory")
        return target

    async def _save(
        self, content: bytes, target: Path, context: ToolContext, expected_version: ArtifactVersion,
    ) -> dict:
        from clearact.tools.artifacts import save_artifact

        # Commit without yielding: a cancelled renderer must not leave a background file write.
        return save_artifact(context, target, content, expected_version=expected_version)

    def _verify_target(self, arguments: dict, context: ToolContext, target: Path) -> None:
        if self._target(arguments, context) != target:
            raise ToolValidationError("Output path changed during generation; no file was written")


class CreateDocumentTool(_OfficeTool):
    name = "create_document"
    suffix = ".docx"

    def definition(self) -> ToolDefinition:
        text_array = {"type": "array", "items": {"type": "string"}, "maxItems": 200}
        table = {
            "type": "object",
            "properties": {
                "headers": {"type": "array", "items": {"type": "string"}, "minItems": 1, "maxItems": 12},
                "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}, "maxItems": 500},
            },
            "required": ["headers", "rows"],
            "additionalProperties": False,
        }
        return ToolDefinition(
            name=self.name,
            description=(
                "Create or replace a native .docx Word document with A4 layout, Chinese fonts, headings, "
                "bullets, tables, and page numbers. Table rows must match headers. Writes use file snapshots."
            ),
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "title": {"type": "string", "maxLength": 300},
                    "subtitle": {"type": "string", "maxLength": 1000},
                    "sections": {
                        "type": "array",
                        "minItems": 1,
                        "maxItems": 80,
                        "items": {
                            "type": "object",
                            "properties": {
                                "heading": {"type": "string"},
                                "paragraphs": text_array,
                                "bullets": text_array,
                                "table": table,
                            },
                            "additionalProperties": False,
                        },
                    },
                },
                "required": ["path", "title", "sections"],
                "additionalProperties": False,
            },
        )

    async def execute(self, arguments: dict, context: ToolContext, action_id: str) -> ToolResult:
        content = _document_content(arguments)
        target = self._target(arguments, context)
        expected_version = capture_artifact_version(target)
        data = await asyncio.to_thread(_render_document, content)
        self._verify_target(arguments, context, target)
        metadata = await self._save(data, target, context, expected_version)
        metadata.update(
            {
                "path": str(target),
                "size_bytes": len(data),
                "file_kind": "docx",
                "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "section_count": len(content["sections"]),
                "table_count": sum("table" in section for section in content["sections"]),
            }
        )
        return ToolResult(
            action_id=action_id,
            tool_name=self.name,
            ok=True,
            content=f"Created Word document {target} ({len(data)} bytes).",
            metadata=metadata,
        )


class CreateSpreadsheetTool(_OfficeTool):
    name = "create_spreadsheet"
    suffix = ".xlsx"

    def definition(self) -> ToolDefinition:
        return ToolDefinition(
            name=self.name,
            description=(
                "Create or replace a native .xlsx Excel workbook with styled tables, filters, frozen headers, "
                "and optional bar/line charts. Column formats/widths are arrays matching headers; formats accept "
                "general/number/text/integer/decimal/percent/currency/date/datetime or Excel format strings. "
                "date/datetime cells accept ISO strings. Strings beginning '=' are formulas unless the format is text. "
                "Formulas are preserved, not calculated here; Excel recalculates when opened. "
                "Writes use file snapshots."
            ),
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "sheets": {
                        "type": "array",
                        "minItems": 1,
                        "maxItems": 12,
                        "items": {
                            "type": "object",
                            "properties": {
                                "name": {"type": "string", "maxLength": 31},
                                "headers": {
                                    "type": "array",
                                    "items": {"type": "string"},
                                    "minItems": 1,
                                    "maxItems": 50,
                                },
                                "rows": {
                                    "type": "array",
                                    "maxItems": 10_000,
                                    "items": {
                                        "type": "array",
                                        "items": {"type": ["string", "number", "boolean", "null"]},
                                    },
                                },
                                "column_formats": {"type": "array", "items": {"type": "string"}},
                                "column_widths": {"type": "array", "items": {"type": ["number", "null"]}},
                                "charts": {
                                    "type": "array",
                                    "maxItems": 4,
                                    "items": {
                                        "type": "object",
                                        "properties": {
                                            "type": {"type": "string", "enum": ["bar", "line"]},
                                            "title": {"type": "string"},
                                            "category_column": {
                                                "type": "integer",
                                                "minimum": 1,
                                                "description": "1-based column",
                                            },
                                            "value_columns": {
                                                "type": "array",
                                                "items": {"type": "integer", "minimum": 1},
                                            },
                                            "anchor": {"type": "string", "description": "Excel cell, e.g. A15"},
                                        },
                                        "required": ["type", "value_columns"],
                                        "additionalProperties": False,
                                    },
                                },
                            },
                            "required": ["name", "headers", "rows"],
                            "additionalProperties": False,
                        },
                    },
                },
                "required": ["path", "sheets"],
                "additionalProperties": False,
            },
        )

    async def execute(self, arguments: dict, context: ToolContext, action_id: str) -> ToolResult:
        sheets = _spreadsheet_content(arguments)
        target = self._target(arguments, context)
        expected_version = capture_artifact_version(target)
        data = await asyncio.to_thread(_render_spreadsheet, sheets)
        self._verify_target(arguments, context, target)
        metadata = await self._save(data, target, context, expected_version)
        formula_count = sum(
            isinstance(value, str) and value.startswith("=") and fmt != "text"
            for sheet in sheets
            for row in sheet["rows"]
            for value, fmt in zip(row, sheet["column_formats"], strict=True)
        )
        metadata.update(
            {
                "path": str(target),
                "size_bytes": len(data),
                "file_kind": "xlsx",
                "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "sheets": [
                    {"name": sheet["name"], "rows": len(sheet["rows"]), "columns": len(sheet["headers"])}
                    for sheet in sheets
                ],
                "chart_count": sum(len(sheet["charts"]) for sheet in sheets),
                "formula_count": formula_count,
                "formulas_calculated": False,
                "recalculate_on_open": True,
            }
        )
        formula_note = " Formulas will recalculate in Excel when opened." if formula_count else ""
        return ToolResult(
            action_id=action_id,
            tool_name=self.name,
            ok=True,
            content=f"Created Excel workbook {target} ({len(data)} bytes).{formula_note}",
            metadata=metadata,
        )

"""Opt-in real localhost model evaluation with isolated files and per-request measurements."""

from __future__ import annotations

import argparse
import asyncio
import json
import time
from collections import Counter
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlsplit

import httpx

from clearact import __version__, cli
from clearact.context.budget import ContextBudget
from clearact.domain.enums import RiskLevel, RunStatus
from clearact.domain.models import ChatMessage, Run, RunExecutionSettings, UserPolicy
from clearact.runtime.prompting import system_prompt
from clearact.settings import load_settings

CASES = {
    "bare_math": "只回复整数结果：17 乘以 23 等于多少？",
    "greeting": "你好",
    "math": "只回复整数结果：17 乘以 23 等于多少？",
    "file": "创建 hello.txt，内容严格为 ClearAct lightweight test，不要创建其他文件。",
    "edit": "修复 calc.py 的 price_after_discount(price, discount)，discount 是折扣比例，"
            "100 元打八折传入 0.2 应返回 80。保留 README.md 不改，不要增加文件。完成后检查语法，不执行命令。",
    "followup": "继续修改 calc.py：discount 必须在 0 到 1 之间，包含端点，超出范围抛出 ValueError。"
                "只修改 calc.py，保留 README.md 不改。完成后检查语法，不执行命令。",
    "web": "创建单文件 index.html 计数器页面，初始显示 0，有加一和重置按钮。"
           "不要增加其他源文件或文档。完成后用浏览器实际测试加一、再次加一和重置，并断言数字分别为 1、2、0。离线完成。",
    "data": "读取 orders.csv，按 category 合计 amount，将结果写入 summary.json，"
            "键为每个类别及 total，值为数值。只增加这个 JSON 文件，不要修改原始 CSV，离线完成。",
    "office": "根据 orders.csv 生成 sales.xlsx 和 report.docx。Excel 的 Sales 工作表保留 category/amount "
              "两列和三条原始记录，A5 为 TOTAL，B5 使用公式 =SUM(B2:B4)。"
              "Word 报告包含总金额和每类合计。只增加这两个原生文件，保留 CSV 不变，离线完成。",
    "context": "读取 input_00.json 到 input_11.json，将每个文件内的 id 和 value 汇总到 summary.json，"
               "id 作键、value 作数值，必须覆盖全部十二个文件，不能猜数值。padding 是无关背景，不用复述。"
               "输入保持不变，只增加 summary.json，离线完成。",
}

CSV_DATA = "category,amount\nA,12.5\nB,7\nA,2.5\n"


def save(path: Path, value) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")


def emit(value) -> None:
    print(json.dumps(value, ensure_ascii=True), flush=True)


def numeric_usage(value):
    if not isinstance(value, dict):
        return {}
    return {key: numeric_usage(item) if isinstance(item, dict) else item for key, item in value.items()
            if isinstance(item, int | float | dict) and not isinstance(item, bool)}


def reported_tokens(requests):
    successful = [request for request in requests if 200 <= (request.get("status") or 0) < 300]

    def total(name, native):
        values = [request.get("usage", {}).get(name, request.get(native)) for request in successful]
        return sum(values) if values and all(type(value) is int for value in values) else None

    prompt, completion = total("prompt_tokens", "prompt_eval_count"), total("completion_tokens", "eval_count")
    return {"prompt_tokens": prompt, "completion_tokens": completion,
            "usage_reported": prompt is not None and completion is not None}


class MeasuredProvider:
    def __init__(self, provider, records, budget):
        self.provider, self.records, self.budget = provider, records, budget
        self.request_timeout_seconds = provider.request_timeout_seconds

    async def chat(self, messages, tools):
        record = {"call": len(self.records) + 1, "messages": len(messages),
                  "tool_schemas": [tool.name for tool in tools],
                  "estimated_input_units": self.budget.estimate(messages, tools),
                  "reasoning_characters_sent": sum(len(message.reasoning_content or "") for message in messages)}
        self.records.append(record)
        emit({"event": "request", **record})
        started = time.monotonic()
        try:
            response = await self.provider.chat(messages, tools)
        except BaseException as exc:
            record.update(elapsed_seconds=round(time.monotonic() - started, 3), error_type=type(exc).__name__)
            raise
        record.update(elapsed_seconds=round(time.monotonic() - started, 3), usage=response.usage,
                      returned_reasoning_characters=len(response.reasoning_content or ""),
                      returned_content_characters=len(response.content or ""),
                      tools_called=[action.tool_name for action in response.tool_calls],
                      finish_reason=response.finish_reason)
        emit({"event": "response", **record})
        return response


class PreviewOnlyApproval:
    async def request(self, action, assessment):
        return action.tool_name == "preview_project" and assessment.category == "browser_interaction"


async def verify(case, project, run, final):
    if case in {"math", "bare_math"}:
        return {"correct_answer": final.strip() == "391"}
    if case == "greeting":
        return {"nonempty_reply": bool(final.strip()), "no_tools": not any(m.tool_calls for m in run.messages)}
    if case == "file":
        path = project / "hello.txt"
        return {"exact_file": path.is_file() and path.read_text(encoding="utf-8").strip()
                == "ClearAct lightweight test",
                "no_extra_files": sorted(p.name for p in project.iterdir()) == ["hello.txt"]}
    if case in {"edit", "followup"}:
        # Inspect the small generated calculator before executing isolated behavioral assertions.
        import ast
        source = (project / "calc.py").read_text(encoding="utf-8")
        tree = ast.parse(source)
        simple_module = all(isinstance(node, ast.FunctionDef) or isinstance(node, ast.Expr)
                            and isinstance(node.value, ast.Constant) for node in tree.body)
        if not simple_module:
            return {"simple_calculator_module": False}
        namespace = {"__builtins__": {"ValueError": ValueError, "TypeError": TypeError,
                                     "isinstance": isinstance, "float": float, "int": int, "round": round}}
        exec(compile(tree, "evaluation-calc.py", "exec"), namespace)
        function = namespace.get("price_after_discount")
        result = {"discount_20_percent": function(100, .2) == 80,
                  "zero_discount": function(100, 0) == 100, "full_discount": function(100, 1) == 0,
                  "readme_preserved": (project / "README.md").read_bytes() == b"Keep this note unchanged.\n",
                  "no_extra_files": {p.name for p in project.iterdir()} == {"calc.py", "README.md"}}
        if case == "followup":
            for value in (-.1, 1.1):
                try:
                    function(100, value)
                except ValueError:
                    result[f"reject_{value}"] = True
                else:
                    result[f"reject_{value}"] = False
        return result
    if case == "web":
        import os
        import re

        from playwright.async_api import async_playwright

        async with async_playwright() as playwright:
            options = {"headless": True}
            if os.name == "nt":
                options["channel"] = "msedge"
            async with await playwright.chromium.launch(**options) as browser:
                page = await browser.new_page()
                await page.route("**/*", lambda route: route.continue_() if route.request.url.startswith("file:")
                                 else route.abort())
                await page.goto((project / "index.html").as_uri())
                checks = {"initial_zero": await page.get_by_text("0", exact=True).is_visible()}
                button = page.get_by_role("button", name=re.compile(r"加一|增加|\+1|increment|add", re.I))
                await button.click(timeout=3000)
                checks["increment_one"] = await page.get_by_text("1", exact=True).is_visible()
                await button.click(timeout=3000)
                checks["increment_two"] = await page.get_by_text("2", exact=True).is_visible()
                await page.get_by_role("button", name=re.compile(r"重置|reset", re.I)).click(timeout=3000)
                checks["reset_zero"] = await page.get_by_text("0", exact=True).is_visible()
        return {**checks, "html_exists": (project / "index.html").is_file(),
                "no_extra_source_files": {p.name for p in project.iterdir() if p.name != ".clearact"} == {"index.html"},
                "model_browser_assertions": sum(m.metadata.get("assertions", 0) for m in run.messages
                                                 if m.name == "preview_project" and
                                                 m.metadata.get("verification_passed")) >= 3}
    if case == "data":
        return {"correct_totals": json.loads((project / "summary.json").read_text(encoding="utf-8"))
                == {"A": 15, "B": 7, "total": 22},
                "source_unchanged": (project / "orders.csv").read_bytes() == CSV_DATA.encode(),
                "no_extra_files": {p.name for p in project.iterdir()} == {"orders.csv", "summary.json"}}
    if case == "office":
        import re

        from docx import Document
        from openpyxl import load_workbook

        workbook = load_workbook(project / "sales.xlsx", data_only=False)
        try:
            sheet = workbook["Sales"]
            rows = list(sheet.iter_rows(min_row=1, max_row=5, min_col=1, max_col=2, values_only=True))
        finally:
            workbook.close()
        document = Document(project / "report.docx")
        text = "\n".join([paragraph.text for paragraph in document.paragraphs] +
                         [cell.text for table in document.tables for row in table.rows for cell in row.cells])
        return {"raw_excel_rows": rows[:4] == [("category", "amount"), ("A", 12.5), ("B", 7), ("A", 2.5)],
                "native_excel_formula": rows[4] == ("TOTAL", "=SUM(B2:B4)"),
                "word_has_totals": all(re.search(r"(?<!\d)" + value + r"(?:\.0+)?(?!\d)", text)
                                       for value in ("22", "15", "7")),
                "source_unchanged": (project / "orders.csv").read_bytes() == CSV_DATA.encode(),
                "no_extra_files": {p.name for p in project.iterdir()} == {"orders.csv", "sales.xlsx", "report.docx"}}
    if case == "context":
        expected = {f"entry_{index:02}": 1007 + index * 37 for index in range(12)}
        sources = [project / f"input_{index:02}.json" for index in range(12)]
        return {"all_values": json.loads((project / "summary.json").read_text(encoding="utf-8")) == expected,
                "sources_unchanged": all(path.read_bytes() == (json.dumps({
                    "id": f"entry_{index:02}", "value": 1007 + index * 37, "padding": "background " * 280,
                }, ensure_ascii=False, indent=2) + "\n").encode("utf-8") for index, path in enumerate(sources)),
                "no_extra_files": {path.name for path in project.iterdir()}
                == {path.name for path in sources} | {"summary.json"}}
    return {}


async def evaluate(output: Path, selected_cases: list[str], timeout: float):
    cli._load_dotenv(cli._project_root())
    settings = load_settings(cli._project_root())
    selected = settings.models["default_profile"]
    profile = settings.models["profiles"][selected]
    if urlsplit(profile["base_url"]).hostname not in {"localhost", "127.0.0.1", "::1"}:
        raise ValueError("This evaluation only calls the configured localhost model")
    if output.exists() and any(output.iterdir()):
        raise ValueError("Output directory must be empty")
    output.mkdir(parents=True, exist_ok=True)
    provider = cli._build_provider(profile)
    report = {"version": __version__, "model": profile["model"], "provider": profile["provider"],
              "reasoning_history": profile.get("reasoning_history", "auto"), "case_timeout_seconds": timeout,
              "context_window": profile["context_window"], "context_budget_ratio": settings.agent.context_budget_ratio,
              "limits": "Single sequential run per case; no sampling override; timings include local service.",
              "cases": []}
    previous = None
    original_client = httpx.AsyncClient
    http_records = []

    class ObservedClient(original_client):
        async def post(self, url, **kwargs):
            record = {"status": None}
            http_records.append(record)
            started = time.monotonic()
            try:
                response = await super().post(url, **kwargs)
                record.update(status=response.status_code)
                try:
                    data = response.json()
                    record["usage"] = numeric_usage(data.get("usage", {}))
                    for key in ("eval_count", "prompt_eval_count", "total_duration", "load_duration",
                                "prompt_eval_duration", "eval_duration"):
                        if isinstance(data.get(key), int | float):
                            record[key] = data[key]
                except ValueError:
                    pass
                return response
            finally:
                record["elapsed_seconds"] = round(time.monotonic() - started, 3)

    for case in selected_cases:
        case_root = output / case
        case_root.mkdir()
        if case == "followup":
            if previous is None:
                raise ValueError("followup needs an earlier edit case")
            run, project, local_settings = previous
            run.messages.append(ChatMessage(role="user", content=CASES[case]))
            run.status = RunStatus.CREATED
        else:
            project = case_root / "project"
            project.mkdir()
            if case == "edit":
                (project / "calc.py").write_text("def price_after_discount(price, discount):\n"
                                                "    return price * (1 + discount)\n", encoding="utf-8")
                (project / "README.md").write_text("Keep this note unchanged.\n", encoding="utf-8", newline="\n")
            if case in {"data", "office"}:
                (project / "orders.csv").write_text(CSV_DATA, encoding="utf-8", newline="\n")
            if case == "context":
                for index in range(12):
                    save(project / f"input_{index:02}.json", {"id": f"entry_{index:02}",
                         "value": 1007 + index * 37, "padding": "background " * 280})
            local_settings = settings.model_copy(update={"workspace_root": project,
                                                         "data_root": case_root / "state", "mcp_servers": {}})
            run = Run(goal=CASES[case], execution=RunExecutionSettings(workdir=str(project), profile=selected),
                      policy=UserPolicy(autonomy_threshold=RiskLevel.YELLOW, allow_web=False,
                                        allowed_scopes=[str(project)], capability_rules={
                                            "local_read": "allow", "workspace_create": "allow",
                                            "workspace_modify": "allow", "web_read": "deny", "mcp_read": "deny",
                                            "mcp_write": "deny", "outside_write": "deny",
                                            "command_execute": "deny", "browser_interaction": "allow",
                                        }), messages=[ChatMessage(role="system", content=system_prompt(CASES[case])),
                                                      ChatMessage(role="user", content=CASES[case])])
        records = []
        measured = MeasuredProvider(provider, records, ContextBudget(
            profile["context_window"], settings.agent.context_budget_ratio,
        ))
        http_start, message_start = len(http_records), len(run.messages)
        started, error, final = time.monotonic(), None, ""
        emit({"event": "case_started", "case": case})
        try:
            with (patch.object(cli, "_build_provider", return_value=measured),
                  patch.object(httpx, "AsyncClient", ObservedClient)):
                if case == "bare_math":
                    response = await asyncio.wait_for(
                        measured.chat([ChatMessage(role="user", content=CASES[case])], []), timeout=timeout,
                    )
                    final = response.content or ""
                else:
                    await asyncio.wait_for(cli._run_impl(
                        run.goal, selected, None, str(project), 16, 40, run=run, render_console=False,
                        approval_gate=PreviewOnlyApproval(), settings=local_settings,
                    ), timeout=timeout)
                    final = run.messages[-1].content or ""
        except Exception as exc:
            error = type(exc).__name__
        runtime_elapsed = round(time.monotonic() - started, 3)
        verification_started = time.monotonic()
        try:
            checks = await verify(case, project, run, final)
        except Exception as exc:
            checks = {"verification_error": type(exc).__name__}
        tools = Counter(message.name for message in run.messages[message_start:] if message.role == "tool")
        result = {"case": case, "prompt": CASES[case], "run_id": run.id,
                  "status": run.status.value if case != "bare_math" else "direct_model",
                  "task_mode": run.task_contract.mode if run.task_contract else None,
                  "delivery_passed": run.delivery_report.passed if run.delivery_report else None,
                  "elapsed_seconds": runtime_elapsed,
                  "verification_seconds": round(time.monotonic() - verification_started, 3), "error_type": error,
                  "model_calls": len(records), "http_requests": http_records[http_start:],
                  **reported_tokens(http_records[http_start:]),
                  "failed_tools": sum(message.role == "tool" and message.metadata.get("status") == "failed"
                                      for message in run.messages[message_start:]),
                  "tools": dict(tools), "checks": checks, "final_answer": final, "calls": records,
                  "passed": error is None and bool(checks) and all(value is True for value in checks.values())
                  and (case == "bare_math" or run.status is RunStatus.COMPLETED)}
        report["cases"].append(result)
        save(case_root / "summary.json", result)
        save(output / "summary.json", report)
        emit({"event": "case_completed", **{key: value for key, value in result.items()
                                              if key not in {"prompt", "calls", "http_requests", "final_answer"}}})
        if case == "edit":
            previous = run, project, local_settings
        # Cancellation may leave local backend generation running. Do not queue more work behind it.
        if error in {"TimeoutError", "ModelRequestError", "ModelServiceUnavailableError"}:
            break
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--cases", nargs="+", choices=list(CASES), default=list(CASES))
    parser.add_argument("--case-timeout", type=float, default=240)
    args = parser.parse_args()
    asyncio.run(evaluate(args.output.resolve(), args.cases, args.case_timeout))
